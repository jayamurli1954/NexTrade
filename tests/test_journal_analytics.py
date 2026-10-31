"""Tests for journal analytics + Momentum Breakdown paper gate."""

from __future__ import annotations

import os
import tempfile
import unittest

from trading_app.excel_trade_logger import ExcelTradeLogger
from trading_app.journal_analytics import (
    allow_momentum_breakdown_signal,
    compute_journal_metrics,
    filter_trades_for_day,
)
from trading_app.opportunity_engine import OpportunityEngine


SAMPLE_TRADES = [
    {
        "trade_id": "t1",
        "timestamp": "2026-07-28 10:00:00",
        "symbol": "AAA",
        "style": "INTRADAY",
        "strategy": "Golden Ratio Scan",
        "direction": "BUY",
        "quantity": 10,
        "entry_price": 100.0,
        "exit_price": 110.0,
        "exit_reason": "AUTO_TARGET_HIT",
        "pnl": 100.0,
        "status": "CLOSED",
    },
    {
        "trade_id": "t2",
        "timestamp": "2026-07-31 11:00:00",
        "symbol": "BBB",
        "style": "INTRADAY",
        "strategy": "Momentum Breakdown",
        "direction": "SELL",
        "quantity": 10,
        "entry_price": 200.0,
        "exit_price": 220.0,
        "exit_reason": "AUTO_SQUARE_OFF_15:15",
        "pnl": -200.0,
        "status": "CLOSED",
    },
    {
        "trade_id": "t3",
        "timestamp": "2026-08-04 09:30:00",
        "symbol": "CCC",
        "style": "INTRADAY",
        "strategy": "Oversold Bounce",
        "direction": "BUY",
        "quantity": 5,
        "entry_price": 50.0,
        "exit_price": 55.0,
        "exit_reason": "DAY_PROFIT_TARGET",
        "pnl": 25.0,
        "status": "CLOSED",
    },
]


class TestJournalAnalytics(unittest.TestCase):
    def test_period_metrics_and_drawdown(self):
        m = compute_journal_metrics(SAMPLE_TRADES)
        self.assertEqual(m.total_trades, 3)
        self.assertEqual(m.winning_trades, 2)
        self.assertEqual(m.gross_pnl, -75.0)
        self.assertEqual(m.period_start, "2026-07-28")
        self.assertEqual(m.period_end, "2026-08-04")
        self.assertLess(m.max_drawdown, 0.0)
        self.assertIn("Momentum Breakdown", m.by_strategy)
        self.assertEqual(m.by_strategy["Momentum Breakdown"]["pnl"], -200.0)

    def test_today_filter(self):
        today = filter_trades_for_day(SAMPLE_TRADES, "2026-08-04")
        self.assertEqual(len(today), 1)
        self.assertEqual(compute_journal_metrics(today).gross_pnl, 25.0)

    def test_mb_gate(self):
        ok, _ = allow_momentum_breakdown_signal(
            {"confidence": 0.8, "rsi": 55},
            {"india_vix": 14.0, "nifty_change_pct": 0.1},
        )
        self.assertTrue(ok)
        bad, reason = allow_momentum_breakdown_signal(
            {"confidence": 0.6, "rsi": 55},
            {"india_vix": 14.0, "nifty_change_pct": 0.1},
        )
        self.assertFalse(bad)
        self.assertIn("0.70", reason)
        bad2, _ = allow_momentum_breakdown_signal(
            {"confidence": 0.8, "rsi": 40},
            {"india_vix": 14.0},
        )
        self.assertFalse(bad2)
        bad3, _ = allow_momentum_breakdown_signal(
            {"confidence": 0.8, "rsi": 55},
            {"india_vix": 11.0},
        )
        self.assertFalse(bad3)

    def test_opportunity_engine_gates_sell(self):
        eng = OpportunityEngine()
        signals = [
            {
                "symbol": "RELIANCE",
                "action": "SELL",
                "confidence": 0.60,
                "current_price": 1000,
                "stop_loss": 1020,
                "target": 960,
                "rsi": 40,
                "reason": "weak",
            }
        ]
        live = {"RELIANCE": {"ltp": 1000.0, "source": "ANGEL_WEBSOCKET"}}
        out = eng.get_top_opportunities(
            live_ltp_map=live,
            signals=signals,
            market_context={"india_vix": 11.0, "nifty_change_pct": 0.5},
        )
        self.assertEqual(out, [])

    def test_excel_report_sheets(self):
        td = tempfile.mkdtemp()
        try:
            logger = ExcelTradeLogger(reports_dir=td)
            path = logger.generate_excel_report(executed_trades=SAMPLE_TRADES, starting_capital=1_000_000)
            self.assertTrue(os.path.exists(path))
            import openpyxl

            wb = openpyxl.load_workbook(path)
            try:
                self.assertIn("Executed Trades", wb.sheetnames)
                self.assertIn("Today vs Period Summary", wb.sheetnames)
                self.assertIn("Daily Breakdown", wb.sheetnames)
                self.assertIn("Strategy Breakdown", wb.sheetnames)
                ws2 = wb["Today vs Period Summary"]
                labels = [ws2.cell(row=r, column=1).value for r in range(1, 40)]
                self.assertTrue(any(x and "Period Start" in str(x) for x in labels))
                self.assertTrue(any(x and "Max Drawdown" in str(x) for x in labels))
                self.assertTrue(any(x and "Expectancy" in str(x) for x in labels))
            finally:
                wb.close()
        finally:
            import shutil

            shutil.rmtree(td, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
